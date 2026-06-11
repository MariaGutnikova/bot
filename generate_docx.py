import re
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def add_toc(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r = run._r
    r.append(fldChar)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)

def set_gost_style(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for i in range(1, 4):
        h_style = doc.styles[f'Heading {i}']
        h_font = h_style.font
        h_font.name = 'Times New Roman'
        h_font.size = Pt(14)
        h_font.bold = True
        h_font.color.rgb = None
        h_paragraph_format = h_style.paragraph_format
        h_paragraph_format.space_before = Pt(14)
        h_paragraph_format.space_after = Pt(14)
        h_paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        h_paragraph_format.first_line_indent = Cm(1.25)

def main():
    with open('/Users/maria/.gemini/antigravity/brain/1cfe2b36-56c6-411c-9b9a-30ff9c8e24ad/diploma_expanded.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    set_gost_style(doc)

    doc.add_heading('ОГЛАВЛЕНИЕ', level=1)
    add_toc(doc)
    doc.add_page_break()

    in_mermaid = False
    mermaid_counter = 0
    diagrams = [
        '/Users/maria/Desktop/1_AS_IS_process.png',
        '/Users/maria/Desktop/2_TO_BE_process.png',
        '/Users/maria/Desktop/3_Architecture.png',
        '/Users/maria/Desktop/4_ER_Diagram.png'
    ]

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        if not line:
            i += 1
            continue

        if line.startswith('```mermaid'):
            in_mermaid = True
            i += 1
            continue
        
        if in_mermaid and line.startswith('```'):
            in_mermaid = False
            if mermaid_counter < len(diagrams):
                img_path = diagrams[mermaid_counter]
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    run = p_img.add_run()
                    run.add_picture(img_path, width=Inches(6.0))
                except Exception as e:
                    print(f"Failed to load image {img_path}: {e}")
                mermaid_counter += 1
            i += 1
            continue
            
        if in_mermaid:
            i += 1
            continue

        # Match Table block
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Filter out separator
            data_rows = []
            for tline in table_lines:
                if set(tline.replace('|', '').replace('-', '').replace(':', '').strip()) == set():
                    continue # separator line
                cells = [cell.strip() for cell in tline.strip('|').split('|')]
                data_rows.append(cells)
                
            if data_rows:
                table = doc.add_table(rows=len(data_rows), cols=len(data_rows[0]))
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(data_rows):
                    for c_idx, cell_text in enumerate(row_data):
                        if c_idx < len(table.columns):
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_text
                            for paragraph in cell.paragraphs:
                                paragraph.paragraph_format.first_line_indent = 0
                                paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                                for run in paragraph.runs:
                                    run.font.name = 'Times New Roman'
                                    run.font.size = Pt(12)
            continue

        # Match Image standard markdown
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = p_img.add_run()
                run.add_picture(img_path, width=Inches(6.0))
            except Exception as e:
                pass
            
            p_cap = doc.add_paragraph(caption)
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.first_line_indent = 0
            i += 1
            continue

        # Match Chapter headings
        if re.match(r'^[А-ЯЁ0-9\s.,-]+$', line) and len(line) > 3 and not line.startswith('Приложение'):
            h = doc.add_heading(line, level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.first_line_indent = 0
            i += 1
            continue

        # Match Subheadings
        m2 = re.match(r'^(\d+\.\d+\.)\s+(.+)$', line)
        if m2:
            doc.add_heading(line, level=2)
            i += 1
            continue
            
        m3 = re.match(r'^(\d+\.\d+\.\d+\.)\s+(.+)$', line)
        if m3:
            doc.add_heading(line, level=3)
            i += 1
            continue

        # Normal text
        p = doc.add_paragraph(line)
        p.paragraph_format.first_line_indent = Cm(1.25)
        
        i += 1

    doc.save('/Users/maria/Desktop/bot/Дипломная_работа_МСБ.docx')

if __name__ == '__main__':
    main()
