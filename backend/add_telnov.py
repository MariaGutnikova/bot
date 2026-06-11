# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document('../Гутникова ВКР_ФИНАЛ.docx')

in_bib = False
bib_start_idx = -1
app_start_idx = -1

for i, p in enumerate(doc.paragraphs):
    if 'СПИСОК' in p.text.upper() and 'ЛИТЕРАТУР' in p.text.upper():
        in_bib = True
        bib_start_idx = i
        continue
        
    if in_bib and p.text.strip().startswith('ПРИЛОЖЕНИЕ'):
        app_start_idx = i
        break

if bib_start_idx != -1 and app_start_idx != -1:
    # Extract current sources
    current_sources = []
    for i in range(bib_start_idx + 1, app_start_idx):
        text = doc.paragraphs[i].text.strip()
        if text:
            # Strip the leading number "1. ", "2. ", etc.
            if ". " in text:
                text = text.split(". ", 1)[1]
            current_sources.append(text)
            
    # Insert Telnov at index 4 (so it becomes item 5)
    telnov = "Тельнов, Ю. Ф. Методы и модели обоснования сценариев применения цифровых двойников бизнес-процессов сетевых предприятий // Бизнес-информатика. – 2023. – Т. 17, № 4. – С. 73-93. – DOI: 10.17323/2587-814X.2023.4.73.93."
    current_sources.insert(4, telnov)
    
    # Delete old paragraphs
    for i in range(app_start_idx - 1, bib_start_idx, -1):
        p = doc.paragraphs[i]
        p._element.getparent().remove(p._element)
        
    # Re-insert with proper numbering
    bib_header = doc.paragraphs[bib_start_idx]
    for idx, source_text in enumerate(reversed(current_sources)):
        num = len(current_sources) - idx
        full_text = f"{num}. {source_text}"
        new_p = bib_header.insert_paragraph_before(full_text)
        bib_header._p.addnext(new_p._p)
        
        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        new_p.paragraph_format.first_line_indent = Inches(0.49)
        new_p.paragraph_format.line_spacing = 1.5
        for run in new_p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)

doc.save('../Гутникова ВКР_ФИНАЛ.docx')
print("Telnov restored. Total sources:", len(current_sources))
