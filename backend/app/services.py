import io
import asyncio
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

def generate_excel_sync(tasks, title="Отчёт по задачам"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    headers = ['ID', 'Задача', 'ID Автора', 'ID Исполнителя', 'Срок', 'Статус', 'Создана', 'Выполнена']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    
    for row, t in enumerate(tasks, 2):
        record = [t.id, t.text, t.author_id, t.assignee_id, t.deadline, t.status, 
                  t.created_at.strftime("%Y-%m-%d %H:%M") if t.created_at else "", 
                  t.completed_at.strftime("%Y-%m-%d %H:%M") if t.completed_at else ""]
        for col, value in enumerate(record, 1):
            cell = ws.cell(row=row, column=col, value=str(value) if value else '—')
            cell.border = thin_border
    
    for col in range(1, 9):
        ws.column_dimensions[chr(64 + col)].width = 20
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

async def generate_excel_report(tasks):
    # Выполняем синхронный код в отдельном потоке, чтобы не блокировать цикл событий
    return await asyncio.to_thread(generate_excel_sync, tasks)

def generate_pie_chart_sync(data_dict, title="Статистика"):
    if not data_dict or sum(data_dict.values()) == 0:
        return None
        
    labels = []
    sizes = []
    colors = []
    
    status_map = {
        'new': ('В работе', '#FF9800'), # Современный оранжевый
        'done': ('Выполнено', '#4CAF50') # Современный зеленый
    }
    
    for status, count in data_dict.items():
        label, color = status_map.get(status, (status, '#9E9E9E'))
        labels.append(f'{label} ({count})')
        sizes.append(count)
        colors.append(color)
        
    fig, ax = plt.subplots(figsize=(8, 6), facecolor='#FAFAFA')
    wedges, texts, autotexts = ax.pie(sizes, labels=labels, colors=colors,
                                        autopct='%1.1f%%', startangle=90,
                                        wedgeprops=dict(width=0.4, edgecolor='w')) # Donut chart стиль
    
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_fontweight('bold')
        
    plt.title(title, fontsize=16, fontweight='bold', pad=20, color='#333333')
    plt.axis('equal')
    
    buf = io.BytesIO()
    plt.savefig(buf, format='png', dpi=150, bbox_inches='tight', facecolor='#FAFAFA')
    buf.seek(0)
    plt.close(fig)
    return buf.read()

async def generate_status_chart(data_dict):
    return await asyncio.to_thread(generate_pie_chart_sync, data_dict)
