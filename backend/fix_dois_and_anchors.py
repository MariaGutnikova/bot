# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = docx.Document('../Гутникова ВКР_ФИНАЛ.docx')

def insert_text_after(paragraph, text):
    new_p = paragraph.insert_paragraph_before(text)
    paragraph._p.addnext(new_p._p)
    new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    new_p.paragraph_format.first_line_indent = Inches(0.49)
    new_p.paragraph_format.line_spacing = 1.5
    for run in new_p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
    return new_p

# 1. Add BI Anchor (Chapter 2.2)
for p in doc.paragraphs:
    if 'Долговечность гарантируется записью упреждающего журнала транзакций' in p.text:
        anchor1 = "С экономической точки зрения, архитектура базы данных и реляционная модель спроектированы с учетом перспективных требований бизнес-аналитики (Business Intelligence). Строгая третья нормальная форма (3NF) позволяет в будущем легко интегрировать систему с корпоративными BI-дашбордами. Это обеспечит руководство компании прозрачной управленческой отчетностью и инструментами для расчета юнит-экономики каждого сотрудника в реальном времени, превращая технический инструмент в полноценный актив для принятия стратегических бизнес-решений."
        insert_text_after(p, anchor1)
        break

# 2. Add BYOD Anchor (Chapter 2.4)
for p in doc.paragraphs:
    if 'Это радикально снижает нагрузку на процессор мобильного устройства и обеспечивает плавную частоту кадров' in p.text:
        anchor2 = "Выбор подобного оптимизированного стека имеет прямое экономическое обоснование. Снижение требований к аппаратной части позволяет компании успешно реализовать корпоративную концепцию BYOD (Bring Your Own Device). Поскольку приложение работает плавно даже на устаревших личных смартфонах сотрудников, у отдела отпадает необходимость в закупке дорогостоящего парка служебных устройств. Это дополнительно снижает капитальные затраты (CAPEX) на аппаратное обеспечение и ускоряет окупаемость всего проекта."
        insert_text_after(p, anchor2)
        break

# 3. Purge DOI and EDN from Bibliography
in_bib = False
for p in doc.paragraphs:
    if 'СПИСОК' in p.text.upper() and 'ЛИТЕРАТУР' in p.text.upper():
        in_bib = True
    
    if in_bib and p.text.strip().startswith('ПРИЛОЖЕНИЕ'):
        in_bib = False
        break
        
    if in_bib and p.text.strip():
        # Match " - DOI: 10...." or " – DOI: 10..."
        original_text = p.text
        # Remove DOI
        new_text = re.sub(r'\s*[\-–]\s*DOI:\s*[^\s]+', '', original_text)
        # Remove EDN
        new_text = re.sub(r'\s*[\-–]\s*EDN:\s*[^\s]+', '', new_text)
        
        # Cleanup trailing dots
        new_text = new_text.strip()
        if new_text.endswith('.'):
            pass # normal
        else:
            new_text += '.'
            
        new_text = new_text.replace('..', '.')
        new_text = new_text.replace('. .', '.')
        
        if new_text != original_text:
            p.text = new_text
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.49)
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)

doc.save('../Гутникова ВКР_ФИНАЛ.docx')
print("Anchors injected and DOIs purged.")
